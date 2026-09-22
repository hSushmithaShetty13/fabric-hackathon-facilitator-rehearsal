from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "presentations"
OUTPUT_PATH = OUTPUT_DIR / "fabric_architecture_three_solutions.pptx"

W = Inches(13.333)
H = Inches(7.5)

INK = RGBColor(28, 37, 44)
MUTED = RGBColor(94, 105, 112)
PAPER = RGBColor(247, 246, 242)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 122, 128)
CORAL = RGBColor(232, 91, 70)
YELLOW = RGBColor(244, 190, 52)
GREEN = RGBColor(67, 145, 94)
BLUE = RGBColor(50, 105, 168)
PALE_TEAL = RGBColor(221, 240, 239)
PALE_CORAL = RGBColor(250, 228, 223)
PALE_YELLOW = RGBColor(251, 241, 207)
PALE_BLUE = RGBColor(225, 235, 247)
GRID = RGBColor(218, 219, 214)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=15, color=INK, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(spacing)
        paragraph.text = "- " + paragraph.text
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_node(slide, title, subtitle, x, y, w, h, fill=WHITE, accent=TEAL):
    add_rect(slide, x, y, w, h, fill, GRID, True)
    add_rect(slide, x, y, 0.08, h, accent, accent)
    add_text(slide, title, x + 0.22, y + 0.15, w - 0.35, 0.34, 15, INK, True)
    add_text(slide, subtitle, x + 0.22, y + 0.53, w - 0.35, h - 0.62, 10.5, MUTED)


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.25):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_pill(slide, text, x, y, w, fill, color=INK):
    add_rect(slide, x, y, w, 0.34, fill, fill, True)
    add_text(slide, text, x, y + 0.01, w, 0.28, 10, color, True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)


def add_header(slide, kicker, title, number):
    add_rect(slide, 0, 0, 13.333, 0.12, TEAL, TEAL)
    add_text(slide, kicker.upper(), 0.58, 0.35, 4.5, 0.28, 10, TEAL, True)
    add_text(slide, title, 0.58, 0.70, 11.8, 0.64, 27, INK, True, "Aptos Display")
    add_text(slide, f"{number:02d}", 12.15, 0.42, 0.6, 0.3, 10, MUTED, True,
             align=PP_ALIGN.RIGHT)


def add_footer(slide):
    add_text(slide, "Microsoft Fabric | Manchester Data Pipelines Hack", 0.58, 7.13, 7.0, 0.2,
             8.5, MUTED)


def set_background(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_metric(slide, value, label, x, y, w, accent):
    add_rect(slide, x, y, w, 0.95, WHITE, GRID, True)
    add_text(slide, value, x + 0.18, y + 0.12, w - 0.36, 0.36, 23, accent, True)
    add_text(slide, label, x + 0.18, y + 0.55, w - 0.36, 0.22, 9.5, MUTED, True)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_rect(slide, 0, 0, 0.22, 7.5, TEAL, TEAL)
    add_rect(slide, 9.65, 0, 3.683, 7.5, INK, INK)
    add_text(slide, "MICROSOFT FABRIC ARCHITECTURE", 0.78, 0.82, 7.8, 0.3, 11, TEAL, True)
    add_text(slide, "Three ways to build\nthe same data product", 0.78, 1.35, 8.2, 1.55, 34, INK, True,
             "Aptos Display")
    add_text(slide, "Simple  |  Solution 2  |  Solution 3", 0.78, 3.20, 7.8, 0.4, 18, MUTED, True)
    add_text(slide, "From rapid rehearsal to governed, multi-engine medallion delivery", 0.78, 3.78,
             7.8, 0.7, 16, INK)
    for index, (label, color, y) in enumerate([
        ("1  FASTEST PATH", YELLOW, 1.15),
        ("2  METADATA-DRIVEN", CORAL, 2.25),
        ("3  MULTI-ENGINE", TEAL, 3.35),
    ]):
        add_rect(slide, 10.18, y, 2.55, 0.76, color, color, True)
        add_text(slide, label, 10.40, y + 0.23, 2.1, 0.28, 12, INK if index < 2 else WHITE,
                 True)
    add_text(slide, "Validated in Hackathon Demo", 10.18, 5.47, 2.55, 0.55, 13, WHITE, True)
    add_text(slide, "Architecture and decision guide", 10.18, 6.02, 2.55, 0.45, 10.5,
             RGBColor(196, 205, 210))
    add_text(slide, "22 September 2026", 0.78, 6.72, 4.0, 0.3, 10, MUTED)


def overview_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Architecture spectrum", "One outcome, three levels of engineering", 2)
    columns = [
        ("SIMPLE", "Manual + notebook", YELLOW, PALE_YELLOW,
         ["Fastest to start", "Fewest moving parts", "Best for rehearsal and exploration"]),
        ("SOLUTION 2", "Control-driven ingestion", CORAL, PALE_CORAL,
         ["Incremental Delta merge", "Reusable metadata pattern", "Isolated Solution2 schema"]),
        ("SOLUTION 3", "Multi-engine medallion", TEAL, PALE_TEAL,
         ["Dedicated Bronze/Silver/Gold", "Pipeline + Dataflow + T-SQL", "Warehouse audit and DQ evidence"]),
    ]
    for index, (name, subtitle, accent, pale, bullets) in enumerate(columns):
        x = 0.58 + index * 4.22
        add_rect(slide, x, 1.62, 3.75, 4.72, WHITE, GRID, True)
        add_rect(slide, x, 1.62, 3.75, 0.16, accent, accent)
        add_text(slide, name, x + 0.28, 2.02, 3.15, 0.35, 18, accent, True)
        add_text(slide, subtitle, x + 0.28, 2.48, 3.15, 0.32, 13, INK, True)
        add_rect(slide, x + 0.28, 3.03, 3.18, 0.76, pale, pale, True)
        add_text(slide, ["Speed", "Repeatability", "Governance"][index], x + 0.45, 3.20,
                 2.85, 0.25, 13, INK, True)
        add_bullets(slide, bullets, x + 0.28, 4.10, 3.15, 1.65, 13)
        add_text(slide, ["Hours", "Days", "Production pattern"][index], x + 0.28, 5.82,
                 3.15, 0.3, 11, MUTED, True)
    add_footer(slide)


def simple_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Solution 1", "Simple: quickest route from CSV to insight", 3)
    add_pill(slide, "BEST FOR: REHEARSAL / PROTOTYPE", 0.58, 1.42, 2.85, PALE_YELLOW)
    nodes = [
        ("CSV files", "Local rehearsal datasets", 0.65, 2.25, 2.0, YELLOW),
        ("Lakehouse Files", "Manual upload", 3.05, 2.25, 2.0, BLUE),
        ("Notebook or Dataflow", "Clean, type, flag DQ", 5.45, 2.25, 2.25, CORAL),
        ("Gold tables", "Facts + dimensions", 8.10, 2.25, 2.0, GREEN),
        ("Power BI", "Model + intervention report", 10.50, 2.25, 2.15, TEAL),
    ]
    for title, subtitle, x, y, w, accent in nodes:
        add_node(slide, title, subtitle, x, y, w, 1.15, WHITE, accent)
    for left, right in zip(nodes, nodes[1:]):
        add_arrow(slide, left[2] + left[4], 2.82, right[2], 2.82)
    add_rect(slide, 0.65, 4.05, 5.72, 2.15, WHITE, GRID, True)
    add_text(slide, "What it optimizes", 0.95, 4.32, 2.4, 0.3, 16, INK, True)
    add_bullets(slide, ["Low setup effort and rapid iteration", "Easy inspection in a single notebook",
                        "Direct path to the stakeholder report"], 0.95, 4.80, 4.95, 1.15, 13)
    add_rect(slide, 6.73, 4.05, 5.92, 2.15, WHITE, GRID, True)
    add_text(slide, "Trade-offs", 7.03, 4.32, 2.4, 0.3, 16, INK, True)
    add_bullets(slide, ["Manual landing and rerun discipline", "Limited operational audit trail",
                        "Harder to scale across many entities"], 7.03, 4.80, 5.0, 1.15, 13)
    add_footer(slide)


def solution2_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Solution 2", "Metadata-driven incremental ingestion", 4)
    add_pill(slide, "BEST FOR: REUSABLE INGESTION", 0.58, 1.42, 2.65, PALE_CORAL)
    add_node(slide, "GitHub", "Seven source CSVs", 0.65, 2.05, 1.75, 1.05, WHITE, BLUE)
    add_node(slide, "load_control.csv", "Enabled, target, mode, keys, watermark", 0.65, 3.55,
             2.20, 1.25, PALE_CORAL, CORAL)
    add_node(slide, "Lookup + Filter", "Read enabled control rows", 3.30, 2.55, 2.05, 1.15,
             WHITE, CORAL)
    add_node(slide, "ForEach", "Sequential per entity", 5.78, 2.55, 1.75, 1.15, WHITE, YELLOW)
    add_node(slide, "Copy activity", "Land Files/landing/<entity>", 8.00, 1.88, 2.15, 1.15,
             WHITE, BLUE)
    add_node(slide, "Parameterized notebook", "Overwrite or Delta merge", 8.00, 3.35, 2.15, 1.15,
             WHITE, CORAL)
    add_node(slide, "lh_station_ops", "Solution2 schema + ingestion audit", 10.58, 2.55, 2.15,
             1.35, PALE_TEAL, TEAL)
    add_arrow(slide, 2.40, 2.57, 3.30, 2.93)
    add_arrow(slide, 2.85, 4.00, 3.30, 3.30)
    add_arrow(slide, 5.35, 3.12, 5.78, 3.12)
    add_arrow(slide, 7.53, 2.93, 8.00, 2.45)
    add_arrow(slide, 9.08, 3.03, 9.08, 3.35)
    add_arrow(slide, 10.15, 3.93, 10.58, 3.30)
    add_rect(slide, 0.65, 5.38, 12.08, 1.15, WHITE, GRID, True)
    add_text(slide, "Incremental logic", 0.92, 5.62, 1.65, 0.28, 13, CORAL, True)
    add_text(slide, "TRUNCATE for dimensions | INCREMENTAL merge for facts | Optional target-max watermark",
             2.72, 5.60, 7.05, 0.32, 12.5, INK, True)
    add_text(slide, "Isolation", 10.12, 5.62, 0.8, 0.28, 12, MUTED, True)
    add_text(slide, "Solution2.*", 11.00, 5.60, 1.35, 0.32, 13, TEAL, True)
    add_footer(slide)


def solution3_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Solution 3", "Multi-engine medallion with operational evidence", 5)
    add_pill(slide, "BEST FOR: GOVERNED DEMO / PRODUCTION PATTERN", 0.58, 1.42, 3.85, PALE_TEAL)
    layers = [
        ("SOURCE", "GitHub CSV + control manifest", 0.58, BLUE, PALE_BLUE),
        ("BRONZE", "Pipeline Copy + Spark notebook\nlh_git_bronze", 3.10, YELLOW, PALE_YELLOW),
        ("SILVER", "Dataflow Gen2 / Power Query\nlh_dataflow_silver", 5.82, CORAL, PALE_CORAL),
        ("GOLD", "T-SQL stored procedure\nwh_storedproc_gold", 8.54, TEAL, PALE_TEAL),
        ("CONSUME", "Semantic model + Power BI", 11.06, GREEN, RGBColor(226, 240, 228)),
    ]
    for name, detail, x, accent, pale in layers:
        add_rect(slide, x, 2.10, 1.98, 2.32, pale, accent, True)
        add_text(slide, name, x + 0.18, 2.37, 1.62, 0.28, 14, accent, True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.16, 3.02, 1.66, 0.85, 11, INK, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for index in range(len(layers) - 1):
        add_arrow(slide, layers[index][2] + 1.98, 3.25, layers[index + 1][2], 3.25, MUTED, 2.5)
    add_rect(slide, 1.58, 4.95, 10.15, 1.37, INK, INK, True)
    add_text(slide, "WAREHOUSE OPERATIONS PLANE", 1.88, 5.18, 2.55, 0.26, 11, YELLOW, True)
    add_text(slide, "pipeline_run", 4.55, 5.18, 1.35, 0.25, 11, WHITE, True)
    add_text(slide, "pipeline_step", 6.15, 5.18, 1.35, 0.25, 11, WHITE, True)
    add_text(slide, "reconciliation", 7.75, 5.18, 1.35, 0.25, 11, WHITE, True)
    add_text(slide, "data quality", 9.35, 5.18, 1.35, 0.25, 11, WHITE, True)
    add_text(slide, "Bad records remain visible in Silver and are excluded from Gold",
             1.88, 5.72, 9.55, 0.28, 11.5, RGBColor(206, 215, 219))
    add_footer(slide)


def comparison_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Comparison", "Choose the engineering level that matches the goal", 6)
    headers = ["Dimension", "Simple", "Solution 2", "Solution 3"]
    rows = [
        ("Primary goal", "Speed", "Repeatable ingest", "Governed delivery"),
        ("Landing", "Manual upload", "Dynamic Pipeline Copy", "Dynamic Pipeline Copy"),
        ("Transform", "Notebook / Dataflow", "Parameterized Spark", "Spark + Dataflow + T-SQL"),
        ("Incremental", "Manual rerun", "Keys + watermark merge", "Bronze merge; deterministic Silver/Gold"),
        ("Isolation", "Chosen Lakehouse", "Solution2 schema", "Dedicated items per layer"),
        ("Monitoring", "Notebook output", "Ingestion audit", "Run, step, reconcile, DQ tables"),
        ("Best fit", "Workshop", "Reusable ingestion", "Teaching + production pattern"),
    ]
    x_positions = [0.58, 3.15, 6.18, 9.40]
    widths = [2.45, 2.90, 3.10, 3.35]
    for index, header in enumerate(headers):
        color = INK if index == 0 else [YELLOW, CORAL, TEAL][index - 1]
        text_color = WHITE if index in (0, 3) else INK
        add_rect(slide, x_positions[index], 1.58, widths[index], 0.55, color, color)
        add_text(slide, header, x_positions[index] + 0.12, 1.72, widths[index] - 0.24, 0.24,
                 12.5, text_color, True)
    for row_index, row in enumerate(rows):
        y = 2.13 + row_index * 0.66
        fill = WHITE if row_index % 2 == 0 else RGBColor(240, 240, 236)
        for col_index, value in enumerate(row):
            add_rect(slide, x_positions[col_index], y, widths[col_index], 0.66, fill, GRID)
            add_text(slide, value, x_positions[col_index] + 0.12, y + 0.14,
                     widths[col_index] - 0.24, 0.34, 10.8, INK, col_index == 0,
                     valign=MSO_ANCHOR.MIDDLE)
    add_footer(slide)


def evidence_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Runtime evidence", "Solution 3 completed end to end", 7)
    add_text(slide, "Validated master run", 0.58, 1.55, 3.0, 0.35, 16, INK, True)
    add_text(slide, "c87ee117-8bd0-4c77-9dd5-220669ece567", 0.58, 1.98, 4.8, 0.3,
             11.5, MUTED)
    add_metric(slide, "COMPLETED", "MASTER PIPELINE STATUS", 0.58, 2.55, 2.65, GREEN)
    add_metric(slide, "7 / 7", "PIPELINE STEPS COMPLETED", 3.48, 2.55, 2.65, TEAL)
    add_metric(slide, "7 / 7", "RECONCILIATIONS BALANCED", 6.38, 2.55, 2.65, BLUE)
    add_metric(slide, "169", "EXPECTED REJECTED ROWS", 9.28, 2.55, 2.65, CORAL)
    add_rect(slide, 0.58, 4.05, 7.15, 2.18, WHITE, GRID, True)
    add_text(slide, "Data quality findings", 0.88, 4.32, 2.7, 0.3, 16, INK, True)
    add_bullets(slide, ["166 reversed assistance timestamps", "1 fulfilled count above request count",
                        "1 negative passenger count", "1 missing station reference"],
                0.88, 4.78, 6.15, 1.15, 12.5)
    add_rect(slide, 8.03, 4.05, 4.70, 2.18, INK, INK, True)
    add_text(slide, "Design principle", 8.35, 4.32, 3.9, 0.3, 14, YELLOW, True)
    add_text(slide, "A successful run is not enough.", 8.35, 4.84, 3.75, 0.36, 17,
             WHITE, True)
    add_text(slide, "Prove row balance, quality outcomes, and traceability for every entity.",
             8.35, 5.35, 3.65, 0.62, 11.5, RGBColor(206, 215, 219))
    add_footer(slide)


def decision_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Decision guide", "Start simple, add controls when the risk demands them", 8)
    prompts = [
        ("Need an answer today?", "Use Simple", "Manual landing + notebook", YELLOW, PALE_YELLOW),
        ("Loading many files repeatedly?", "Use Solution 2", "Control table + incremental merge", CORAL, PALE_CORAL),
        ("Need governance and operational proof?", "Use Solution 3", "Medallion + audit + DQ", TEAL, PALE_TEAL),
    ]
    for index, (question, answer, detail, accent, pale) in enumerate(prompts):
        y = 1.62 + index * 1.52
        add_rect(slide, 0.58, y, 12.15, 1.16, WHITE, GRID, True)
        add_rect(slide, 0.58, y, 0.18, 1.16, accent, accent)
        add_text(slide, question, 0.98, y + 0.24, 4.55, 0.34, 16, INK, True)
        add_rect(slide, 5.72, y + 0.20, 2.38, 0.76, pale, pale, True)
        add_text(slide, answer, 5.90, y + 0.40, 2.02, 0.25, 13, accent, True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, detail, 8.50, y + 0.38, 3.75, 0.35, 13, MUTED, True)
    add_rect(slide, 0.58, 6.35, 12.15, 0.48, INK, INK, True)
    add_text(slide, "The three approaches are complementary: use them to teach progression, not to declare one universal winner.",
             0.88, 6.47, 11.55, 0.24, 11.5, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide)


def build_deck():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Microsoft Fabric Architecture: Three Solutions"
    prs.core_properties.subject = "Simple, metadata-driven, and multi-engine Fabric architectures"
    prs.core_properties.author = "Manchester Data Pipelines Hack"
    title_slide(prs)
    overview_slide(prs)
    simple_slide(prs)
    solution2_slide(prs)
    solution3_slide(prs)
    comparison_slide(prs)
    evidence_slide(prs)
    decision_slide(prs)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_deck())